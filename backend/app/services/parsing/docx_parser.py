"""Word (.docx) 文本抽取。"""
from docx import Document


def parse_docx(path: str) -> str:
    doc = Document(path)
    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # 表格内容也纳入解析
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)
