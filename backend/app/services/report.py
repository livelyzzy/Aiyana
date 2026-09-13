"""报表生成与导出（Excel / PDF，可嵌入图表）。

使用统一的「报表载荷」结构：
{
  "title": str,
  "meta": [(标签, 值), ...],
  "sections": [{"name": str, "headers": [str], "rows": [[...], ...]}, ...],
  "charts": [{"title": str, "type": "radar"|"bar", "labels": [...], "values": [...]}, ...]
}
"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Image as RLImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.charts import bar_chart, radar_chart

pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
_PDF_FONT = "STSong-Light"


def _render_chart(chart: dict) -> bytes:
    if chart.get("type") == "bar":
        return bar_chart(chart["labels"], chart["values"], chart["title"])
    return radar_chart(chart["labels"], chart["values"], chart["title"])


# ---------------- Excel ----------------
def to_excel(payload: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "评价报告"

    ws.merge_cells("A1:E1")
    ws["A1"] = payload["title"]
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    row = 3
    for label, value in payload.get("meta", []):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    header_fill = PatternFill("solid", fgColor="4C78A8")
    for section in payload.get("sections", []):
        row += 1
        ws.cell(row=row, column=1, value=section["name"]).font = Font(bold=True, size=12)
        row += 1
        headers = section.get("headers", [])
        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=col, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = header_fill
        for r in section.get("rows", []):
            row += 1
            for col, v in enumerate(r, start=1):
                ws.cell(row=row, column=col, value=v)
        row += 1

    # 图表作为图片插入
    img_row = row + 1
    for i, chart in enumerate(payload.get("charts", [])):
        png = _render_chart(chart)
        img = XLImage(BytesIO(png))
        img.width = 360 if chart.get("type") == "bar" else 280
        img.height = 200 if chart.get("type") == "bar" else 280
        ws.add_image(img, f"A{img_row}")
        img_row += 16

    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 24

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------- PDF ----------------
def to_pdf(payload: dict) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=15 * mm, bottomMargin=15 * mm)
    story: list = []

    story.append(Paragraph(payload["title"], style=_h1()))
    story.append(Spacer(1, 6 * mm))

    for label, value in payload.get("meta", []):
        story.append(Paragraph(f"<b>{label}</b>：{value}", style=_body()))
    story.append(Spacer(1, 6 * mm))

    for section in payload.get("sections", []):
        story.append(Paragraph(section["name"], style=_h2()))
        headers = section.get("headers", [])
        rows = section.get("rows", [])
        if headers:
            data = [headers] + rows
            table = Table(data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, 0), _PDF_FONT),
                        ("FONTNAME", (0, 1), (-1, -1), _PDF_FONT),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("GRID", (0, 0), (-1, -1), 0.4, "#999999"),
                        ("BACKGROUND", (0, 0), (-1, 0), "#4C78A8"),
                        ("TEXTCOLOR", (0, 0), (-1, 0), "#FFFFFF"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            story.append(table)
        story.append(Spacer(1, 5 * mm))

    for chart in payload.get("charts", []):
        png = _render_chart(chart)
        img = RLImage(BytesIO(png), width=120 * mm, height=90 * mm)
        story.append(img)
        story.append(Spacer(1, 4 * mm))

    story.append(Spacer(1, 6 * mm))
    story.append(
        Paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", style=_body())
    )

    doc.build(story)
    return buf.getvalue()


def _h1():
    from reportlab.lib.styles import ParagraphStyle

    return ParagraphStyle("h1", fontName=_PDF_FONT, fontSize=18, leading=24, spaceAfter=6)


def _h2():
    from reportlab.lib.styles import ParagraphStyle

    return ParagraphStyle("h2", fontName=_PDF_FONT, fontSize=13, leading=18, spaceBefore=8, spaceAfter=4)


def _body():
    from reportlab.lib.styles import ParagraphStyle

    return ParagraphStyle("body", fontName=_PDF_FONT, fontSize=10, leading=14)
